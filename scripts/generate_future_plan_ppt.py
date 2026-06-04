from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path
from typing import List

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
SOURCE_MD = ROOT / "建设方案" / "01_未来规划.md"
OUTPUT_PPTX = ROOT / "建设方案" / "未来规划方向_汇报版.pptx"

THEME_NAVY = RGBColor(18, 42, 76)
THEME_BLUE = RGBColor(36, 92, 171)
THEME_TEXT = RGBColor(34, 34, 34)
THEME_MUTED = RGBColor(96, 109, 128)
THEME_BG = RGBColor(247, 249, 252)
THEME_ACCENT = RGBColor(232, 240, 254)
THEME_SKY = RGBColor(222, 235, 255)
THEME_GREEN = RGBColor(224, 244, 236)
THEME_GOLD = RGBColor(255, 245, 214)
THEME_RED = RGBColor(253, 235, 235)
WHITE = RGBColor(255, 255, 255)


@dataclass
class Line:
    kind: str
    text: str


@dataclass
class Section:
    title: str
    lines: List[Line] = field(default_factory=list)


def parse_markdown(path: Path) -> tuple[str, List[Section]]:
    title = "未来规划方向"
    sections: List[Section] = []
    current: Section | None = None

    for raw in path.read_text(encoding="utf-8").splitlines():
        line = raw.strip()
        if not line:
            continue

        if line.startswith("# "):
            title = line[2:].strip()
            continue

        if line.startswith("## "):
            current = Section(title=line[3:].strip())
            sections.append(current)
            continue

        if current is None:
            continue

        if line.startswith("### "):
            current.lines.append(Line("subheading", line[4:].strip()))
        elif line.startswith("- "):
            current.lines.append(Line("bullet", line[2:].strip()))
        else:
            current.lines.append(Line("text", line))

    return title, sections


def split_section_title(title: str) -> tuple[str, str]:
    if "：" in title:
        return tuple(title.split("：", 1))
    return "", title


def group_lines(lines: List[Line]) -> List[tuple[str, List[str]]]:
    groups: List[tuple[str, List[str]]] = []
    current_title = ""
    current_lines: List[str] = []

    def flush() -> None:
        nonlocal current_title, current_lines
        if current_lines:
            groups.append((current_title, current_lines))
            current_title = ""
            current_lines = []

    for line in lines:
        if line.kind == "subheading":
            flush()
            current_title = line.text
        else:
            current_lines.append(line.text)
    flush()
    return groups


def add_background(slide) -> None:
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = THEME_BG

    top_band = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.28))
    top_band.fill.solid()
    top_band.fill.fore_color.rgb = THEME_NAVY
    top_band.line.fill.background()

    side_band = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(0.25), Inches(7.5))
    side_band.fill.solid()
    side_band.fill.fore_color.rgb = THEME_BLUE
    side_band.line.fill.background()

    accent = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, Inches(10.8), Inches(0.45), Inches(2.4), Inches(2.4))
    accent.fill.solid()
    accent.fill.fore_color.rgb = THEME_ACCENT
    accent.line.fill.background()


def add_label(slide, text: str, x: float, y: float, w: float = 1.2, h: float = 0.38) -> None:
    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    box.fill.solid()
    box.fill.fore_color.rgb = THEME_NAVY
    box.line.fill.background()
    tf = box.text_frame
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = text
    run.font.name = "PingFang SC"
    run.font.size = Pt(10)
    run.font.bold = True
    run.font.color.rgb = WHITE


def add_title_block(slide, title: str) -> str:
    page_label, main_title = split_section_title(title)
    if page_label:
        add_label(slide, page_label, 0.78, 0.5, 1.5, 0.42)

    box = slide.shapes.add_textbox(Inches(0.8), Inches(1.0), Inches(10.8), Inches(0.65))
    tf = box.text_frame
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = main_title
    run.font.name = "PingFang SC"
    run.font.size = Pt(24)
    run.font.bold = True
    run.font.color.rgb = THEME_NAVY

    rule = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.82), Inches(1.62), Inches(1.4), Inches(0.08))
    rule.fill.solid()
    rule.fill.fore_color.rgb = THEME_BLUE
    rule.line.fill.background()
    return main_title


def add_textbox(slide, x: float, y: float, w: float, h: float, text: str, size: int, color: RGBColor, bold: bool = False, align: PP_ALIGN | None = None) -> None:
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    if align is not None:
        p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = "PingFang SC"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def add_card(slide, x: float, y: float, w: float, h: float, title: str | None = None, body_lines: List[str] | None = None, fill: RGBColor = WHITE, title_color: RGBColor = THEME_NAVY, body_color: RGBColor = THEME_TEXT, bullet: bool = False, body_size: int = 14, title_size: int = 16) -> None:
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = RGBColor(225, 231, 239)

    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.18)
    tf.margin_right = Inches(0.18)
    tf.margin_top = Inches(0.12)
    tf.margin_bottom = Inches(0.1)

    first_para = True
    if title:
        p = tf.paragraphs[0]
        run = p.add_run()
        run.text = title
        run.font.name = "PingFang SC"
        run.font.size = Pt(title_size)
        run.font.bold = True
        run.font.color.rgb = title_color
        p.space_after = Pt(7)
        first_para = False

    if body_lines:
        for line in body_lines:
            p = tf.paragraphs[0] if first_para else tf.add_paragraph()
            first_para = False
            p.space_after = Pt(5)
            if bullet:
                p.bullet = True
                p.level = 0
            run = p.add_run()
            run.text = line
            run.font.name = "PingFang SC"
            run.font.size = Pt(body_size)
            run.font.color.rgb = body_color


def add_step_card(slide, x: float, y: float, w: float, h: float, step_no: int, text: str) -> None:
    add_card(slide, x, y, w, h, fill=WHITE)
    badge = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x + 0.18), Inches(y + 0.16), Inches(0.48), Inches(0.48))
    badge.fill.solid()
    badge.fill.fore_color.rgb = THEME_BLUE
    badge.line.fill.background()
    tf = badge.text_frame
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = f"{step_no:02d}"
    run.font.name = "PingFang SC"
    run.font.size = Pt(12)
    run.font.bold = True
    run.font.color.rgb = WHITE
    add_textbox(slide, x + 0.18, y + 0.78, w - 0.32, h - 0.95, text, 13, THEME_TEXT)


def add_cover_slide(prs: Presentation, deck_title: str) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_background(slide)

    banner = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.7), Inches(0.72), Inches(11.2), Inches(0.38))
    banner.fill.solid()
    banner.fill.fore_color.rgb = THEME_BLUE
    banner.line.fill.background()

    hero = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.75), Inches(1.32), Inches(7.4), Inches(2.0))
    hero.fill.solid()
    hero.fill.fore_color.rgb = WHITE
    hero.line.color.rgb = RGBColor(225, 231, 239)

    title_box = slide.shapes.add_textbox(Inches(1.05), Inches(1.65), Inches(6.6), Inches(1.0))
    title_frame = title_box.text_frame
    title_frame.word_wrap = True
    p = title_frame.paragraphs[0]
    run = p.add_run()
    run.text = deck_title
    run.font.name = "PingFang SC"
    run.font.size = Pt(28)
    run.font.bold = True
    run.font.color.rgb = THEME_NAVY

    subtitle_box = slide.shapes.add_textbox(Inches(1.06), Inches(2.45), Inches(6.4), Inches(0.8))
    subtitle_frame = subtitle_box.text_frame
    subtitle_frame.word_wrap = True
    p = subtitle_frame.paragraphs[0]
    run = p.add_run()
    run.text = "LightRAG 第一阶段价值、数据链建设与长期演进路径"
    run.font.name = "PingFang SC"
    run.font.size = Pt(18)
    run.font.color.rgb = THEME_MUTED

    right_panel = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.5), Inches(1.32), Inches(3.95), Inches(3.65))
    right_panel.fill.solid()
    right_panel.fill.fore_color.rgb = THEME_NAVY
    right_panel.line.fill.background()
    add_textbox(slide, 8.85, 1.72, 3.2, 0.35, "汇报主题", 12, WHITE, True)
    add_textbox(slide, 8.85, 2.15, 3.0, 1.8, "第一阶段价值\n数据链建设\n长期演进路径", 24, WHITE, True)
    add_textbox(slide, 8.85, 4.18, 2.9, 0.45, "GraphRAG + MCP + LLM", 14, THEME_SKY, True)

    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(4.4), Inches(11.4), Inches(1.55))
    card.fill.solid()
    card.fill.fore_color.rgb = THEME_ACCENT
    card.line.color.rgb = THEME_ACCENT

    card_text = slide.shapes.add_textbox(Inches(1.1), Inches(4.82), Inches(10.7), Inches(0.7))
    frame = card_text.text_frame
    frame.word_wrap = True
    p = frame.paragraphs[0]
    run = p.add_run()
    run.text = "先把经验变成资产，再把资产升级成智能能力"
    run.font.name = "PingFang SC"
    run.font.size = Pt(22)
    run.font.bold = True
    run.font.color.rgb = THEME_BLUE
    p.alignment = PP_ALIGN.CENTER


def add_positioning_slide(prs: Presentation, section: Section) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_background(slide)
    add_title_block(slide, section.title)
    bullets = [line.text for line in section.lines if line.kind == "bullet"]
    texts = [line.text for line in section.lines if line.kind == "text"]
    add_card(slide, 0.85, 1.95, 5.15, 4.55, "第一阶段定位", bullets[:3], WHITE, bullet=True, body_size=15)
    add_card(slide, 6.2, 1.95, 6.15, 4.55, "未来系统核心能力", [bullets[3]], THEME_NAVY, WHITE, WHITE, False, 18, 16)
    if texts:
        add_card(slide, 0.85, 6.0, 11.45, 0.82, body_lines=texts, fill=THEME_ACCENT, body_color=THEME_BLUE, body_size=18)


def add_benefits_slide(prs: Presentation, section: Section) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_background(slide)
    add_title_block(slide, section.title)
    bullets = [line.text for line in section.lines if line.kind == "bullet"]
    titles = ["先查知识助手", "缩短定位时间", "沉淀标准知识", "能力可复制", "支持治理可视化"]
    fills = [WHITE, THEME_ACCENT, WHITE, WHITE, THEME_ACCENT]
    positions = [
        (0.85, 1.95, 3.55, 2.05),
        (4.62, 1.95, 3.55, 2.05),
        (8.39, 1.95, 3.55, 2.05),
        (0.85, 4.22, 5.45, 2.1),
        (6.48, 4.22, 5.46, 2.1),
    ]
    for (x, y, w, h), title, body, fill in zip(positions, titles, bullets, fills):
        add_card(slide, x, y, w, h, title, [body], fill=fill, bullet=False, body_size=15)


def add_construction_slide(prs: Presentation, section: Section) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_background(slide)
    add_title_block(slide, section.title)
    bullets = [line.text for line in section.lines if line.kind == "bullet"]
    x_positions = [0.7, 3.22, 5.74, 8.26, 10.78]
    for idx, (x, text) in enumerate(zip(x_positions, bullets), start=1):
        add_step_card(slide, x, 2.1, 2.1, 3.8, idx, text)


def add_data_chain_slide(prs: Presentation, section: Section) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_background(slide)
    add_title_block(slide, section.title)
    groups = group_lines(section.lines)
    sources = groups[0][1] if groups else []
    chain = groups[1][1] if len(groups) > 1 else []
    add_label(slide, "数据来源", 0.86, 1.95, 1.3, 0.36)
    for idx, text in enumerate(sources):
        add_card(slide, 0.85, 2.35 + idx * 1.34, 4.45, 1.12, body_lines=[text], fill=WHITE, body_size=14)
    add_label(slide, "数据链形成方式", 5.65, 1.95, 1.75, 0.36)
    for idx, text in enumerate(chain, start=1):
        add_step_card(slide, 5.65, 2.35 + (idx - 1) * 1.1, 6.1, 0.9, idx, text)


def add_comparison_slide(prs: Presentation, section: Section) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_background(slide)
    add_title_block(slide, section.title)
    bullets = [line.text for line in section.lines if line.kind == "bullet"]
    texts = [line.text for line in section.lines if line.kind == "text"]
    add_card(slide, 0.9, 2.0, 5.6, 4.3, "传统 RAG", bullets[:3], THEME_GOLD, THEME_NAVY, THEME_TEXT, True, 15)
    add_card(slide, 6.85, 2.0, 5.55, 4.3, "图谱知识库", bullets[3:], THEME_SKY, THEME_NAVY, THEME_TEXT, True, 15)
    if texts:
        add_card(slide, 0.9, 6.05, 11.5, 0.8, body_lines=texts, fill=THEME_NAVY, body_color=WHITE, body_size=18)


def add_example_slide(prs: Presentation, section: Section) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_background(slide)
    add_title_block(slide, section.title)
    groups = group_lines(section.lines)
    question = groups[0][1][0] if groups and groups[0][1] else ""
    trad = groups[1][1] if len(groups) > 1 else []
    graph = groups[2][1] if len(groups) > 2 else []
    summary = [line.text for line in section.lines if line.kind == "text"]
    add_card(slide, 0.88, 1.95, 11.45, 1.0, "示例问题", [question], THEME_ACCENT, THEME_BLUE, THEME_NAVY, False, 16)
    add_card(slide, 0.88, 3.25, 5.55, 3.1, "传统 RAG 的处理方式", trad, THEME_GOLD, THEME_NAVY, THEME_TEXT, True, 14)
    add_card(slide, 6.78, 3.25, 5.55, 3.1, "GraphRAG 的处理方式", graph, THEME_SKY, THEME_NAVY, THEME_TEXT, True, 14)
    if summary:
        add_card(slide, 0.88, 6.55, 11.45, 0.72, body_lines=summary, fill=THEME_NAVY, body_color=WHITE, body_size=18)


def add_output_slide(prs: Presentation, section: Section) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_background(slide)
    add_title_block(slide, section.title)
    groups = group_lines(section.lines)
    outputs = groups[0][1] if groups else []
    promises = groups[1][1] if len(groups) > 1 else []
    targets = groups[2][1] if len(groups) > 2 else []
    add_card(slide, 0.85, 1.95, 5.2, 3.6, "第一阶段产出", outputs, WHITE, THEME_NAVY, THEME_TEXT, True, 14)
    add_card(slide, 6.25, 1.95, 6.0, 3.6, "建议承诺结果", promises, THEME_ACCENT, THEME_NAVY, THEME_TEXT, True, 14)
    chip_x = [0.85, 3.95, 7.05, 10.15]
    chip_w = 2.65
    for x, text in zip(chip_x, targets):
        add_card(slide, x, 5.9, chip_w, 1.05, body_lines=[text], fill=THEME_NAVY, body_color=WHITE, body_size=13)


def add_roadmap_slide(prs: Presentation, section: Section) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_background(slide)
    add_title_block(slide, section.title)
    groups = group_lines(section.lines)
    positions = [(0.85, 1.95), (6.7, 1.95), (0.85, 4.25), (6.7, 4.25)]
    fills = [WHITE, THEME_SKY, WHITE, THEME_ACCENT]
    for idx, ((title, lines), (x, y), fill) in enumerate(zip(groups, positions, fills), start=1):
        short_title = title.split("（", 1)[0]
        add_card(slide, x, y, 5.75, 1.9, f"{idx}. {short_title}", lines, fill=fill, bullet=True, body_size=14)


def add_roles_slide(prs: Presentation, section: Section) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_background(slide)
    add_title_block(slide, section.title)
    bullets = [line.text for line in section.lines if line.kind == "bullet"]
    role_bullets = bullets[:5]
    governance = bullets[5:]
    y = 1.95
    role_fills = [WHITE, THEME_SKY, WHITE, THEME_ACCENT, WHITE]
    for text, fill in zip(role_bullets, role_fills):
        add_card(slide, 0.85, y, 5.45, 0.82, body_lines=[text], fill=fill, body_size=13)
        y += 0.95
    add_card(slide, 6.6, 1.95, 5.7, 4.95, "安全与治理原则", governance, THEME_NAVY, WHITE, WHITE, True, 15)


def add_summary_slide(prs: Presentation, section: Section) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_background(slide)
    add_title_block(slide, section.title)
    bullets = [line.text for line in section.lines if line.kind == "bullet"]
    add_label(slide, "核心结论", 0.88, 1.95, 1.3, 0.36)
    add_card(slide, 0.88, 2.35, 11.4, 1.0, body_lines=[bullets[0]], fill=THEME_RED, body_color=THEME_NAVY, body_size=18)
    add_card(slide, 0.88, 3.6, 3.55, 2.05, "本质", [bullets[1]], WHITE, THEME_BLUE, THEME_TEXT, False, 15)
    add_card(slide, 4.58, 3.6, 3.55, 2.05, "直接价值", [bullets[2]], THEME_ACCENT, THEME_BLUE, THEME_TEXT, False, 15)
    add_card(slide, 8.28, 3.6, 3.99, 2.05, "长期价值", [bullets[3]], WHITE, THEME_BLUE, THEME_TEXT, False, 15)
    add_card(slide, 0.88, 5.95, 11.4, 0.95, "最终形态", [bullets[4]], THEME_NAVY, WHITE, WHITE, False, 17)


def add_fallback_slide(prs: Presentation, section: Section) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_background(slide)
    add_title_block(slide, section.title)
    bullets = [line.text for line in section.lines if line.kind == "bullet"]
    texts = [line.text for line in section.lines if line.kind == "text"]
    add_card(slide, 0.9, 1.95, 11.35, 4.85, body_lines=bullets, fill=WHITE, bullet=True, body_size=15)
    if texts:
        add_card(slide, 0.9, 6.15, 11.35, 0.7, body_lines=texts, fill=THEME_ACCENT, body_color=THEME_BLUE, body_size=16)


def add_content_slide(prs: Presentation, section: Section) -> None:
    title = section.title
    if "项目定位" in title:
        add_positioning_slide(prs, section)
    elif "第一阶段能带来什么" in title:
        add_benefits_slide(prs, section)
    elif "第一阶段建设内容" in title:
        add_construction_slide(prs, section)
    elif "数据来源与数据链" in title:
        add_data_chain_slide(prs, section)
    elif "为什么不是传统知识库" in title:
        add_comparison_slide(prs, section)
    elif "对比示例" in title:
        add_example_slide(prs, section)
    elif "产出与可承诺结果" in title:
        add_output_slide(prs, section)
    elif "长期迭代计划" in title:
        add_roadmap_slide(prs, section)
    elif "角色化智能体与安全治理" in title:
        add_roles_slide(prs, section)
    elif "汇报总结" in title:
        add_summary_slide(prs, section)
    else:
        add_fallback_slide(prs, section)


def build_presentation(source_path: Path, output_path: Path) -> Path:
    deck_title, sections = parse_markdown(source_path)

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    add_cover_slide(prs, deck_title)
    for section in sections:
        add_content_slide(prs, section)

    prs.save(output_path)
    return output_path


if __name__ == "__main__":
    output = build_presentation(SOURCE_MD, OUTPUT_PPTX)
    print(output)